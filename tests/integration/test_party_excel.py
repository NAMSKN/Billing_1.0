"""Excel export / import tests for the party master (product rules 15-18)."""

from __future__ import annotations

from datetime import UTC
from pathlib import Path

import pytest
from openpyxl import load_workbook

from invoice_generator.infrastructure.db.connection import connect
from invoice_generator.infrastructure.db.migrator import apply_pending
from tests.support.id_factory import SequentialIdGenerator
from vendor_customer.application.dto import CreatePartyCommand, PartyInput
from vendor_customer.application.party_service import PartyService
from vendor_customer.domain.models import PartyAddress, PartyType
from vendor_customer.infrastructure.db.sqlite_party_repository import SqlitePartyRepository
from vendor_customer.infrastructure.excel import columns as C
from vendor_customer.infrastructure.excel.export_service import (
    ExportScope,
    PartyExcelExporter,
    select_for_scope,
)
from vendor_customer.infrastructure.excel.import_service import (
    DuplicateResolution,
    ImportAs,
    PartyExcelImporter,
)


class _Clock:
    def now(self):  # type: ignore[no-untyped-def]
        from datetime import datetime

        return datetime(2026, 9, 11, tzinfo=UTC)


@pytest.fixture()
def service() -> PartyService:
    conn = connect(":memory:")
    apply_pending(conn)
    return PartyService(
        SqlitePartyRepository(conn), id_generator=SequentialIdGenerator(1), clock=_Clock()
    )


def _india(name: str, ptype: PartyType) -> PartyInput:
    return PartyInput(
        company_name=name,
        company_type=ptype,
        billing_address=PartyAddress(state="Maharashtra", state_code="27", city="Pune"),
    )


def _seed(service: PartyService) -> None:
    service.create_party(CreatePartyCommand(_india("CustOnly", PartyType.CUSTOMER)))
    service.create_party(CreatePartyCommand(_india("VendOnly", PartyType.VENDOR)))
    service.create_party(CreatePartyCommand(_india("BothInc", PartyType.CUSTOMER_VENDOR)))


def _names(path: Path) -> list[str]:
    wb = load_workbook(path)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]
    name_idx = header.index(C.COL_NAME)
    return [r[name_idx] for r in rows[1:]]


def test_select_for_scope_dedup() -> None:
    from vendor_customer.domain.models import Party

    cust = Party(company_name="C", company_type=PartyType.CUSTOMER)
    vend = Party(company_name="V", company_type=PartyType.VENDOR)
    both = Party(company_name="B", company_type=PartyType.CUSTOMER_VENDOR)
    parties = [cust, vend, both]
    assert {p.company_name for p in select_for_scope(parties, ExportScope.CUSTOMERS)} == {"C", "B"}
    assert {p.company_name for p in select_for_scope(parties, ExportScope.VENDORS)} == {"V", "B"}
    got_both = select_for_scope(parties, ExportScope.BOTH)
    assert [p.company_name for p in got_both] == ["C", "V", "B"]  # each once


def test_export_customers(service: PartyService, tmp_path: Path) -> None:
    _seed(service)
    exporter = PartyExcelExporter()
    out = exporter.export(service.list_parties(), ExportScope.CUSTOMERS, tmp_path / "c.xlsx")
    names = _names(out)
    assert "CustOnly" in names
    assert "BothInc" in names  # customer/vendor included in customer export
    assert "VendOnly" not in names


def test_export_vendors_includes_customer_vendor(service: PartyService, tmp_path: Path) -> None:
    _seed(service)
    exporter = PartyExcelExporter()
    out = exporter.export(service.list_parties(), ExportScope.VENDORS, tmp_path / "v.xlsx")
    names = _names(out)
    assert "VendOnly" in names
    assert "BothInc" in names
    assert "CustOnly" not in names


def test_export_both_each_once(service: PartyService, tmp_path: Path) -> None:
    _seed(service)
    exporter = PartyExcelExporter()
    out = exporter.export(service.list_parties(), ExportScope.BOTH, tmp_path / "b.xlsx")
    names = _names(out)
    assert sorted(names) == ["BothInc", "CustOnly", "VendOnly"]
    assert names.count("BothInc") == 1


def test_export_column_order(service: PartyService, tmp_path: Path) -> None:
    _seed(service)
    out = PartyExcelExporter().export(service.list_parties(), ExportScope.BOTH, tmp_path / "o.xlsx")
    wb = load_workbook(out)
    header = list(next(wb.active.iter_rows(values_only=True)))
    assert header == list(C.EXPORT_COLUMNS)
    assert "UUID" not in header and "id" not in [h.lower() for h in header]


def test_export_no_silent_overwrite(service: PartyService, tmp_path: Path) -> None:
    _seed(service)
    target = tmp_path / "dup.xlsx"
    PartyExcelExporter().export(service.list_parties(), ExportScope.BOTH, target)
    with pytest.raises(FileExistsError):
        PartyExcelExporter().export(service.list_parties(), ExportScope.BOTH, target)


# --- import ---


def _write_xlsx(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    wb.save(str(path))


def test_import_as_customer(service: PartyService, tmp_path: Path) -> None:
    path = tmp_path / "imp.xlsx"
    _write_xlsx(
        path,
        [C.COL_NAME, C.COL_STATE, C.COL_CITY],
        [["New Cust", "Maharashtra", "Pune"]],
    )
    importer = PartyExcelImporter(service)
    preview = importer.build_preview(path, import_as=ImportAs.CUSTOMER)
    assert not preview.has_header_errors
    assert len(preview.valid_rows) == 1
    result = importer.commit(preview)
    assert result.created == 1
    saved = list(service.list_parties())
    assert saved[0].company_type is PartyType.CUSTOMER


def test_import_as_vendor_and_both(service: PartyService, tmp_path: Path) -> None:
    path = tmp_path / "iv.xlsx"
    _write_xlsx(path, [C.COL_NAME, C.COL_STATE, C.COL_CITY], [["V1", "Maharashtra", "Pune"]])
    importer = PartyExcelImporter(service)
    importer.commit(importer.build_preview(path, import_as=ImportAs.VENDOR))
    assert list(service.list_parties())[0].company_type is PartyType.VENDOR

    path2 = tmp_path / "ib.xlsx"
    _write_xlsx(path2, [C.COL_NAME, C.COL_STATE, C.COL_CITY], [["B1", "Maharashtra", "Pune"]])
    importer.commit(importer.build_preview(path2, import_as=ImportAs.CUSTOMER_VENDOR))
    both = [p for p in service.list_parties() if p.company_name == "B1"]
    assert both[0].company_type is PartyType.CUSTOMER_VENDOR


def test_import_preserves_file_company_type(service: PartyService, tmp_path: Path) -> None:
    path = tmp_path / "ift.xlsx"
    _write_xlsx(
        path,
        [C.COL_NAME, C.COL_COMPANY_TYPE, C.COL_STATE, C.COL_CITY],
        [["FromFile", "Customer / Vendor", "Maharashtra", "Pune"]],
    )
    importer = PartyExcelImporter(service)
    preview = importer.build_preview(path, import_as=ImportAs.FILE)
    assert preview.valid_rows[0].data.company_type is PartyType.CUSTOMER_VENDOR


def test_import_invalid_row_reported(service: PartyService, tmp_path: Path) -> None:
    path = tmp_path / "err.xlsx"
    _write_xlsx(
        path,
        [C.COL_NAME, C.COL_STATE, C.COL_CITY, C.COL_GSTIN, C.COL_REGISTRATION_TYPE],
        [
            ["", "Maharashtra", "Pune", "", ""],  # missing name
            ["NoState", "", "", "", ""],  # missing India state + city
            ["BadGstin", "Maharashtra", "Pune", "NOTVALID", "Regular"],  # invalid gstin
        ],
    )
    importer = PartyExcelImporter(service)
    preview = importer.build_preview(path, import_as=ImportAs.FILE)
    assert len(preview.error_rows) == 3
    assert any("Company Name" in e for e in preview.rows[0].errors)
    assert any("State is required" in e for e in preview.rows[1].errors)
    assert any("GSTIN" in e for e in preview.rows[2].errors)
    # Malformed rows are not persisted.
    result = importer.commit(preview)
    assert result.created == 0


def test_import_blank_eway_and_additional_accepted(service: PartyService, tmp_path: Path) -> None:
    path = tmp_path / "blank.xlsx"
    _write_xlsx(
        path,
        [C.COL_NAME, C.COL_STATE, C.COL_CITY, C.COL_EWAY_DISTANCE, C.COL_NOTE, C.COL_DUE_DAYS],
        [["Blanks OK", "Maharashtra", "Pune", "", "", ""]],
    )
    importer = PartyExcelImporter(service)
    preview = importer.build_preview(path, import_as=ImportAs.FILE)
    assert len(preview.valid_rows) == 1
    row = preview.valid_rows[0]
    assert row.data.distance_for_eway_bill_km is None
    assert row.data.due_days is None


def test_import_duplicate_detection_and_cancel(service: PartyService, tmp_path: Path) -> None:
    service.create_party(
        CreatePartyCommand(
            PartyInput(
                company_name="Dup Co",
                contact_no="555",
                billing_address=PartyAddress(state="Maharashtra", state_code="27", city="Pune"),
            )
        )
    )
    path = tmp_path / "dup.xlsx"
    _write_xlsx(
        path,
        [C.COL_NAME, C.COL_CONTACT_NO, C.COL_STATE, C.COL_CITY],
        [["Dup Co", "555", "Maharashtra", "Pune"]],
    )
    importer = PartyExcelImporter(service)
    preview = importer.build_preview(path, import_as=ImportAs.FILE)
    assert len(preview.duplicate_rows) == 1
    assert preview.rows[0].resolution is DuplicateResolution.SKIP
    # Cancelling (never committing) mutates nothing: still just the one party.
    assert len(service.list_parties()) == 1
    # Committing with default SKIP resolution does not create a duplicate.
    result = importer.commit(preview)
    assert result.skipped == 1 and result.created == 0
    assert len(service.list_parties()) == 1


def test_import_reopen_after_import(service: PartyService, tmp_path: Path) -> None:
    path = tmp_path / "reopen.xlsx"
    _write_xlsx(path, [C.COL_NAME, C.COL_STATE, C.COL_CITY], [["Reopen Co", "Maharashtra", "Pune"]])
    importer = PartyExcelImporter(service)
    importer.commit(importer.build_preview(path, import_as=ImportAs.CUSTOMER))
    reopened = [p for p in service.list_parties() if p.company_name == "Reopen Co"]
    assert len(reopened) == 1
    fetched = service.get_party(reopened[0].id)
    assert fetched is not None and fetched.company_name == "Reopen Co"
