"""Excel import for the party master (product rules, sections 16-18).

Flow: read the file, validate headers, read rows, validate each row, build a
preview (with row-level errors and duplicate detection), let the user confirm,
then persist only the valid rows. A malformed row is never partially created.
Import is offline; no network is used.

"Import As" lets the user force a company type; otherwise a valid COMPANY TYPE
value from the file is preserved (section 16).
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from enum import StrEnum
from pathlib import Path

from openpyxl import load_workbook

from vendor_customer.application.dto import CreatePartyCommand, PartyInput, UpdatePartyCommand
from vendor_customer.application.party_service import PartyService
from vendor_customer.domain.models import (
    INDIA,
    BalanceType,
    OpeningBalance,
    Party,
    PartyAddress,
    PartyType,
    RegistrationType,
)
from vendor_customer.domain.rules import validate_party
from vendor_customer.infrastructure.excel import columns as C

_TYPE_ALIASES: dict[str, PartyType] = {
    "customer": PartyType.CUSTOMER,
    "vendor": PartyType.VENDOR,
    "customer / vendor": PartyType.CUSTOMER_VENDOR,
    "customer/vendor": PartyType.CUSTOMER_VENDOR,
    "customer-vendor": PartyType.CUSTOMER_VENDOR,
    "customer_vendor": PartyType.CUSTOMER_VENDOR,
    "both": PartyType.CUSTOMER_VENDOR,
}

_REG_ALIASES: dict[str, RegistrationType] = {
    "unregistered": RegistrationType.UNREGISTERED,
    "regular": RegistrationType.REGULAR,
    "regular-sez": RegistrationType.REGULAR_SEZ,
    "regular sez": RegistrationType.REGULAR_SEZ,
    "regular_sez": RegistrationType.REGULAR_SEZ,
    "regular-uin": RegistrationType.REGULAR_UIN,
    "regular uin": RegistrationType.REGULAR_UIN,
    "regular_uin": RegistrationType.REGULAR_UIN,
}

_BALANCE_ALIASES: dict[str, BalanceType] = {
    "debit": BalanceType.DEBIT,
    "dr": BalanceType.DEBIT,
    "credit": BalanceType.CREDIT,
    "cr": BalanceType.CREDIT,
}


class ImportAs(StrEnum):
    """Optional company-type override for the whole import (section 16)."""

    FILE = "FILE"  # preserve each row's COMPANY TYPE value
    CUSTOMER = "CUSTOMER"
    VENDOR = "VENDOR"
    CUSTOMER_VENDOR = "CUSTOMER_VENDOR"


class DuplicateResolution(StrEnum):
    """How to handle a row that matches an existing party (section 18)."""

    UPDATE = "UPDATE"
    SKIP = "SKIP"
    CREATE_NEW = "CREATE_NEW"


@dataclass
class ImportRow:
    """A parsed import row plus its validation / duplicate status."""

    row_number: int
    data: PartyInput | None
    errors: list[str] = field(default_factory=list)
    duplicate_of: Party | None = None
    resolution: DuplicateResolution = DuplicateResolution.CREATE_NEW

    @property
    def is_valid(self) -> bool:
        return not self.errors and self.data is not None

    @property
    def is_duplicate(self) -> bool:
        return self.duplicate_of is not None


@dataclass
class ImportPreview:
    """Result of parsing/validating a file before commit."""

    header_errors: list[str] = field(default_factory=list)
    rows: list[ImportRow] = field(default_factory=list)

    @property
    def has_header_errors(self) -> bool:
        return bool(self.header_errors)

    @property
    def valid_rows(self) -> list[ImportRow]:
        return [r for r in self.rows if r.is_valid]

    @property
    def error_rows(self) -> list[ImportRow]:
        return [r for r in self.rows if not r.is_valid]

    @property
    def duplicate_rows(self) -> list[ImportRow]:
        return [r for r in self.rows if r.is_valid and r.is_duplicate]


@dataclass
class ImportResult:
    """Outcome of committing a preview."""

    created: int = 0
    updated: int = 0
    skipped: int = 0


class PartyExcelImporter:
    """Parses, validates and commits party rows from an ``.xlsx`` file."""

    def __init__(self, service: PartyService) -> None:
        self._service = service

    # --- parse / validate ---

    def build_preview(
        self,
        file_path: str | Path,
        *,
        import_as: ImportAs = ImportAs.FILE,
    ) -> ImportPreview:
        """Read and validate ``file_path`` into a preview (no persistence)."""
        preview = ImportPreview()
        workbook = load_workbook(filename=str(file_path), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows_iter = sheet.iter_rows(values_only=True)
            try:
                header = next(rows_iter)
            except StopIteration:
                preview.header_errors.append("The file is empty.")
                return preview

            header_index = _index_headers(header)
            missing = [h for h in C.REQUIRED_IMPORT_HEADERS if h not in header_index]
            if missing:
                preview.header_errors.append(
                    "Missing required column(s): " + ", ".join(missing)
                )
                return preview

            for offset, raw in enumerate(rows_iter, start=2):
                if _is_blank_row(raw):
                    continue
                preview.rows.append(self._parse_row(offset, raw, header_index, import_as))
        finally:
            workbook.close()
        return preview

    def _parse_row(
        self,
        row_number: int,
        raw: tuple[object, ...],
        header_index: dict[str, int],
        import_as: ImportAs,
    ) -> ImportRow:
        def get(name: str) -> object:
            return _cell(raw, header_index, name)

        errors: list[str] = []
        company_type, type_error = _resolve_company_type(get(C.COL_COMPANY_TYPE), import_as)
        if type_error:
            errors.append(type_error)

        registration, reg_error = _resolve_registration(get(C.COL_REGISTRATION_TYPE))
        if reg_error:
            errors.append(reg_error)

        distance, dist_error = _parse_optional_decimal(
            get(C.COL_EWAY_DISTANCE), "E-Way Bill distance"
        )
        if dist_error:
            errors.append(dist_error)

        # Credit Limit is not among the client's sample import columns.
        credit_limit: Decimal | None = None

        due_days, due_error = _parse_optional_int(get(C.COL_DUE_DAYS), "Due Days")
        if due_error:
            errors.append(due_error)

        cust_balance, cb_error = _parse_balance(
            get(C.COL_CUST_BAL_TYPE), get(C.COL_CUST_BAL_AMOUNT), "Customer Balance"
        )
        errors.extend(cb_error)
        vend_balance, vb_error = _parse_balance(
            get(C.COL_VENDOR_BAL_TYPE), get(C.COL_VENDOR_BAL_AMOUNT), "Vendor Balance"
        )
        errors.extend(vb_error)

        country = _text(get(C.COL_COUNTRY)) or INDIA
        billing = PartyAddress(
            address1=_text(get(C.COL_ADDRESS1)),
            address2=_text(get(C.COL_ADDRESS2)),
            landmark=_text(get(C.COL_LANDMARK)),
            country=country,
            state=_text(get(C.COL_STATE)),
            state_code="",
            city=_text(get(C.COL_CITY)),
            pincode=_text(get(C.COL_PINCODE)),
        )

        data = PartyInput(
            company_name=_text(get(C.COL_NAME)),
            company_type=company_type,
            contact_person=_text(get(C.COL_CONTACT_PERSON)),
            contact_no=_text(get(C.COL_CONTACT_NO)),
            email=_text(get(C.COL_EMAIL)),
            registration_type=registration,
            gstin=_text(get(C.COL_GSTIN)).upper(),
            pan=_text(get(C.COL_PAN)).upper(),
            billing_address=billing,
            distance_for_eway_bill_km=distance,
            credit_limit=credit_limit,
            due_days=due_days,
            bank_name=_text(get(C.COL_BANK_NAME)),
            bank_ifsc_code=_text(get(C.COL_BANK_IFSC)).upper(),
            bank_account_number=_text(get(C.COL_BANK_ACCOUNT)),
            fax_no=_text(get(C.COL_FAX)),
            website=_text(get(C.COL_WEBSITE)),
            note=_text(get(C.COL_NOTE)),
            custom_field_1=_text(get(C.COL_CUSTOM1)),
            custom_field_2=_text(get(C.COL_CUSTOM2)),
            custom_field_3=_text(get(C.COL_CUSTOM3)),
            customer_balance=cust_balance,
            vendor_balance=vend_balance,
        )

        # Domain-rule validation (name required, India state/city, formats).
        errors.extend(validate_party(_input_to_party(data)))

        row = ImportRow(row_number=row_number, data=data if not errors else None, errors=errors)
        if not errors:
            duplicate = self._service.find_duplicate(data)
            if duplicate is not None:
                row.duplicate_of = duplicate
                row.resolution = DuplicateResolution.SKIP
        return row

    # --- commit ---

    def commit(self, preview: ImportPreview) -> ImportResult:
        """Persist valid rows honoring each duplicate resolution.

        Error rows and rows resolved as SKIP are not persisted. A row resolved
        as UPDATE overwrites the matched party; CREATE_NEW inserts a new one.
        """
        result = ImportResult()
        for row in preview.valid_rows:
            assert row.data is not None
            if row.is_duplicate:
                if row.resolution is DuplicateResolution.SKIP:
                    result.skipped += 1
                    continue
                if row.resolution is DuplicateResolution.UPDATE and row.duplicate_of is not None:
                    self._service.update_party(
                        UpdatePartyCommand(id=row.duplicate_of.id, data=row.data, is_active=True)
                    )
                    result.updated += 1
                    continue
            self._service.create_party(CreatePartyCommand(data=row.data))
            result.created += 1
        return result


# --- parsing helpers (module-level, pure) ---


def _index_headers(header: tuple[object, ...]) -> dict[str, int]:
    index: dict[str, int] = {}
    for i, value in enumerate(header):
        if value is None:
            continue
        index[str(value).strip()] = i
    return index


def _cell(raw: tuple[object, ...], header_index: dict[str, int], name: str) -> object:
    idx = header_index.get(name)
    if idx is None or idx >= len(raw):
        return None
    return raw[idx]


def _is_blank_row(raw: tuple[object, ...]) -> bool:
    return all(v is None or str(v).strip() == "" for v in raw)


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _resolve_company_type(value: object, import_as: ImportAs) -> tuple[PartyType, str | None]:
    if import_as is ImportAs.CUSTOMER:
        return PartyType.CUSTOMER, None
    if import_as is ImportAs.VENDOR:
        return PartyType.VENDOR, None
    if import_as is ImportAs.CUSTOMER_VENDOR:
        return PartyType.CUSTOMER_VENDOR, None
    text = _text(value)
    if not text:
        # No file value and no override: default to Customer.
        return PartyType.CUSTOMER, None
    resolved = _TYPE_ALIASES.get(text.lower())
    if resolved is None:
        return PartyType.CUSTOMER, f"Invalid Company Type: {text!r}."
    return resolved, None


def _resolve_registration(value: object) -> tuple[RegistrationType, str | None]:
    text = _text(value)
    if not text:
        return RegistrationType.UNREGISTERED, None
    resolved = _REG_ALIASES.get(text.lower())
    if resolved is None:
        return RegistrationType.UNREGISTERED, f"Invalid Registration Type: {text!r}."
    return resolved, None


def _parse_optional_decimal(value: object, label: str) -> tuple[Decimal | None, str | None]:
    text = _text(value)
    if not text:
        return None, None  # blank is valid (product rules, section 17)
    try:
        return Decimal(text), None
    except (InvalidOperation, ValueError):
        return None, f"Invalid {label}: {text!r}."


def _parse_optional_int(value: object, label: str) -> tuple[int | None, str | None]:
    text = _text(value)
    if not text:
        return None, None
    try:
        return int(Decimal(text)), None
    except (InvalidOperation, ValueError):
        return None, f"Invalid {label}: {text!r}."


def _parse_balance(
    type_value: object, amount_value: object, label: str
) -> tuple[OpeningBalance, list[str]]:
    errors: list[str] = []
    type_text = _text(type_value)
    balance_type = BalanceType.DEBIT
    if type_text:
        resolved = _BALANCE_ALIASES.get(type_text.lower())
        if resolved is None:
            errors.append(f"Invalid {label} Type: {type_text!r}.")
        else:
            balance_type = resolved
    amount, amount_error = _parse_optional_decimal(amount_value, f"{label} Amount")
    if amount_error:
        errors.append(amount_error)
    return OpeningBalance(balance_type=balance_type, amount=amount or Decimal("0")), errors


def _input_to_party(data: PartyInput) -> Party:
    """Assemble a Party for validation without touching persistence."""
    return Party(
        company_type=data.company_type,
        company_name=data.company_name.strip(),
        contact_person=data.contact_person.strip(),
        contact_no=data.contact_no.strip(),
        email=data.email.strip(),
        registration_type=data.registration_type,
        gstin=data.gstin.strip().upper(),
        pan=data.pan.strip().upper(),
        billing_address=data.billing_address,
        shipping_address=data.shipping_address,
        distance_for_eway_bill_km=data.distance_for_eway_bill_km,
        bank_name=data.bank_name.strip(),
        bank_ifsc_code=data.bank_ifsc_code.strip().upper(),
        bank_account_number=data.bank_account_number.strip(),
        fax_no=data.fax_no.strip(),
        website=data.website.strip(),
        credit_limit=data.credit_limit,
        due_days=data.due_days,
        note=data.note.strip(),
        custom_field_1=data.custom_field_1.strip(),
        custom_field_2=data.custom_field_2.strip(),
        custom_field_3=data.custom_field_3.strip(),
        customer_balance=data.customer_balance,
        vendor_balance=data.vendor_balance,
    )


def parties_sequence(preview: ImportPreview) -> Sequence[ImportRow]:
    """Convenience accessor for all parsed rows."""
    return preview.rows
